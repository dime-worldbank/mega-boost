# Databricks notebook source
# MAGIC %run ../utils

# COMMAND ----------

# MAGIC %md
# MAGIC # Bulgaria BOOST expenditures 2020 onward
# MAGIC See README.md for the method, the inputs and outputs, the year rules and the checks. Every year from 2020 whose
# MAGIC extract sits in `TXT/` is processed and written to its own file; a new year needs no change here unless the
# MAGIC Ministry of Finance changes the files' layout.

# COMMAND ----------

import json
import re

import numpy as np
import pandas as pd

COUNTRY = 'Bulgaria'
OUT_DIR = Path(prepare_microdata_csv_dir(COUNTRY))
BASE = Path(f"{RAW_INPUT_DIR}/Bulgaria")
REFERENCE_WORKBOOK = BASE / "Bulgaria BOOST 2015-2024 expenditure.xlsx"  # the delivered data, for the check cell
REFERENCE_SHEET = "2019-24"

# Years and files: every extract "YYYY_annual_deatiled_data*.txt" of 2020 or later in TXT/, with the report
# "YYYY - special_spending_units*.xls[x]" of the same year ("-final" versions included; earlier versions, whose
# name carries "v1", ignored). Exactly one file of each kind per year.
def input_file(year, pattern):
    files = sorted(f for f in (BASE / "TXT").glob(pattern.format(year=year)) if "v1" not in f.name.lower())
    assert len(files) == 1, f"{year}: expected one file {pattern.format(year=year)}, found {[f.name for f in files]}"
    return files[0].name
YEARS = sorted({int(f.name[:4]) for f in (BASE / "TXT").glob("2???_annual_deatiled_data*.txt") if int(f.name[:4]) >= 2020})
INPUTS = {year: (input_file(year, "{year}_annual_deatiled_data*.txt"), input_file(year, "{year} - special_spending_units*.xls*"))
          for year in YEARS}
print("years:", ", ".join(f"{y} ({e}, {r})" for y, (e, r) in INPUTS.items()))
LABELS_PATH = BASE / "labels_en.json"  # the code list shared with the 2005-2019 notebook (a copy sits in the repository)


def applymap(frame, fn):
    """DataFrame.applymap, cell by cell (the method is called map from pandas 2.1 and applymap is gone in 3.0)."""
    return frame.applymap(fn) if hasattr(frame, "applymap") else frame.map(fn)

# COMMAND ----------

# 1. Code lists: the hierarchical code list (labels_en.json) gives every unit its admin1-3 labels, every activity
#    its func1-3, every economic code its econ1-2 and expenditure type, every financing source its fin_source1 label
#    and the fin_source2 the workbooks attach to it, and every programme code its fin_source2 label.
code_list = json.load(open(LABELS_PATH, encoding="utf-8"))
econ1_of = {k: v["econ1"] for k, v in code_list["economic"].items()}
econ2_of = {k: v["econ2"] for k, v in code_list["economic"].items()}
exptype_of_code = {k: v["exp_type"] for k, v in code_list["economic"].items()}
exptype_of = {v["econ2"]: v["exp_type"] for v in code_list["economic"].values()}
fin1_of = {int(k): v["fin_source1"] for k, v in code_list["fin_source1"].items()}
fin2_of = {int(k): v.get("fin_source2", "") for k, v in code_list["fin_source1"].items()}  # by financing source ...
fin2_of.update({int(k): v["fin_source2"] for k, v in code_list["fin_source2"].items() if int(k) >= 98000})  # ... or programme code
admin1_of, admin2_of, admin3_of = ({k: v[c] for k, v in code_list["units"].items()} for c in ("admin1", "admin2", "admin3"))
func1_of, func2_of, func3_of = ({k: v[c] for k, v in code_list["activities"].items()} for c in ("func1", "func2", "func3"))
print(f"code lists: {len(econ2_of)} economic codes, {len(admin3_of)} units, {len(func3_of)} activities, {len(fin1_of)} financing sources")

def prints_codes(text, para_col):
    """A report prints paragraph codes when its best column holds at least 50 (the 2021 report prints none)."""
    return int(text[para_col].str.match(r"^\d{2}-\d{2}").sum()) >= 50
line_para = {}
for year, (_, report_file) in INPUTS.items():
    text = applymap(pd.read_excel(BASE / "TXT" / report_file, sheet_name=0, header=None), lambda v: re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else "")
    name_col = next(j for j in text.columns if text[j].map(lambda v: "EXPENDITURE BY FUNCTION" in v).any())
    para_col = max(text.columns, key=lambda j: int(text[j].str.match(r"^\d{2}-\d{2}").sum()))
    if not prints_codes(text, para_col):
        continue
    block = text.loc[text.index[text[name_col].str.contains("EXPENDITURE BY FUNCTION")][0]:]
    for name, para in zip(block[name_col], block[para_col]):
        if name and re.match(r"^\d{2}-\d{2}", para):
            assert line_para.setdefault(name, para) == para, (name, para, line_para[name])
print(f"report line names with a paragraph: {len(line_para)}")

# COMMAND ----------

EXP_COLS = ["year", "admin1", "admin2", "admin3", "func1", "func2", "func3", "econ1", "econ2",
            "fin_source1", "fin_source2", "exp_type", "transfer", "adjusted", "executed"]
# 2-6. One year at a time
expenditures = {}
for YEAR, (extract_file, report_file) in INPUTS.items():
    print(f"\n===== {YEAR}")
    # 2. Parse the extract
    raw = pd.read_csv(BASE / "TXT" / extract_file, sep="\t", dtype=str, keep_default_na=False, encoding="latin-1")
    raw.columns = [c.strip().lower() for c in raw.columns]
    raw = raw.apply(lambda s: s.str.strip())
    raw = raw[raw["quarter_id"] != ""].reset_index(drop=True)
    raw["unit"] = raw["budget_unit_id"].str.zfill(4)
    sub = raw["sub_para_id"].str.lstrip("0")
    raw["econ_code"] = np.where(sub.str.len() == 3, "0" + sub.str[0] + "." + sub.str[-2:], sub.str[:2] + "." + sub.str[-2:])
    raw["activity"] = raw["activity_id"].str.lstrip("0").str[-3:]
    raw["ibsf"] = pd.to_numeric(raw["ibsf_type_id"]).astype(int)
    raw["act"] = pd.to_numeric(raw["act_type_id"]).astype(int)
    raw["op"] = pd.to_numeric(raw["op_code_id"]).astype(int)
    raw["adjusted"] = pd.to_numeric(raw["adj_budget_amt"], errors="coerce")
    raw["executed"] = pd.to_numeric(raw["actual_amt"], errors="coerce")
    raw["has_amount"] = (raw["adjusted"].fillna(0) != 0) | (raw["executed"].fillna(0) != 0)
    print(f"extract: {len(raw):,} rows, {int(raw['has_amount'].sum()):,} with an amount")

    # 3. Clean the expenditure rows
    clean = raw[(raw["para_type_id"] == "2") & raw["has_amount"]].reset_index(drop=True)
    if YEAR >= 2023:
        paragraphs_with_sub = set(clean.loc[clean["econ_code"].str[-2:] != "00", "econ_code"].str[:2])
        clean["filter"] = (clean["econ_code"].str[-2:] != "00") | ~clean["econ_code"].str[:2].isin(paragraphs_with_sub)
    else:
        if YEAR == 2020:
            clean = raw[raw["has_amount"]].reset_index(drop=True)
        elif YEAR == 2021:
            clean["unit_is_text"] = clean["budget_unit_id"].str.lstrip("0").str.len() == 3
            clean["unit_num"] = pd.to_numeric(clean["budget_unit_id"])
            clean["activity_num"] = pd.to_numeric(clean["activity_id"])
            clean = clean.sort_values(["unit_is_text", "unit_num", "activity_num", "act", "ibsf", "op", "econ_code"], kind="stable").reset_index(drop=True)
        elif YEAR == 2022:
            clean = raw[raw["para_type_id"] == "2"].reset_index(drop=True)
            first_other = int((clean["unit"].map(admin1_of) != "2 Local").idxmax())
            clean = pd.concat([clean.iloc[first_other:], clean.iloc[:first_other]]).reset_index(drop=True)
        code, nxt = clean["econ_code"], clean["econ_code"].shift(-1).fillna("")
        clean["filter"] = (code.str[-2:] != "00") | (code.str[:2] != nxt.str[:2])
        if YEAR == 2021:
            clean["filter"] &= code != "29.90"
        clean = clean[(clean["para_type_id"] == "2") & clean["has_amount"]].reset_index(drop=True)
    print(f"expenditure rows with an amount: {len(clean):,}; the paragraph-total rule keeps {int(clean['filter'].sum()):,}")

    if YEAR == 2020:
        nhif = clean[~clean["filter"] & (clean["econ_code"] == "39.00")]
        clean = pd.concat([clean[clean["filter"]], nhif]).reset_index(drop=True)
        clean["filter"] = True
        clean.loc[clean["econ_code"] == "40.71", "econ_code"] = "57.01"
        print(f"{len(nhif)} rows of 39.00 kept")
    if YEAR == 2022:
        has_sub = clean["econ_code"].str[-2:] != "00"
        sub_keys = set(clean.loc[has_sub, "unit"] + "/" + clean.loc[has_sub, "activity_id"] + "/" + clean.loc[has_sub, "econ_code"].str[:2])
        clean.loc[~(clean["unit"] + "/" + clean["activity_id"] + "/" + clean["econ_code"].str[:2]).isin(sub_keys), "filter"] = True
        clean.loc[clean["econ_code"] == "40.00", "filter"] = False
    clean = clean[clean["filter"]].reset_index(drop=True)

    # 4. Parse the special-units report
    sheet = pd.read_excel(BASE / "TXT" / report_file, sheet_name=0, header=None)
    sheet = applymap(sheet, lambda v: re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v)
    text = applymap(sheet, lambda v: v if isinstance(v, str) else "")
    first_data = text.index[applymap(text, lambda v: "REVENUE" in v).any(axis=1)][0]
    header = {j: " / ".join(v for v in text.loc[:first_data - 1, j] if v) for j in sheet.columns}
    name_col = next(j for j in sheet.columns if text[j].map(lambda v: "EXPENDITURE BY FUNCTION" in v).any())
    ssu_cols = [j for j, h in header.items() if "SSU" in h or ("Special" in h and "Budget" in h)]
    amount = sum(pd.to_numeric(sheet[j], errors="coerce").fillna(0.0) for j in ssu_cols)
    para_col = max(sheet.columns, key=lambda j: int(text[j].str.match(r"^\d{2}-\d{2}").sum()))
    if prints_codes(text, para_col):
        para = text[para_col].where(text[para_col].str.match(r"^\d{2}-\d{2}"), "")
    else:  # 2021: paragraphs from the line names
        para = text[name_col].map(line_para).fillna("")
    next_para = para[para != ""].shift(-1).reindex(para.index).fillna("")
    leaf = (para.str[-2:] != "00") | (para.str[:2] != next_para.str[:2])
    econ = para.str[:2] + "." + (para.str[3:5] if YEAR == 2020 else para.str[-2:])
    print(f"report: {len(sheet)} rows; amount columns {[header[j][:32] for j in ssu_cols]}")

    FUNC2_BY_HEADER = [(r"EXECUTIVE", "011"), (r"GENERAL SERVICES", "012"), (r"SCIENCE", "161"),
                       (r"^A\.?\s*DEFEN", "021"), (r"POLICE", "022"), (r"JUDICIAL", "023"), (r"PRIS", "024"),
                       (r"CIVIL PROTECTION", "025"), (r"PENSIONS", "051"), (r"EMPLOYMENT", "053"),
                       (r"SOCIAL ASSISTANCE", "052"), (r"^A\.?\s*HOUSING", "061"), (r"ENVIRONMENT", "062"),
                       (r"RECREATION", "071"), (r"PHYSICAL", "072"), (r"\bCULTURE", "073"), (r"RELIGIO", "074"),
                       (r"MINING", "081"), (r"AGRICULTURE", "082"), (r"TRANSPORT", "083"), (r"INDUSTRY", "084"),
                       (r"TOURISM", "085"), (r"OTHER ECONOMIC", "086")]
    ROMAN = {"I": "011", "II": "021", "III": "031", "IV": "041", "V": "051", "VI": "061", "VII": "071", "VIII": "081", "IX": "091"}
    start = text.index[text[name_col].str.contains("EXPENDITURE BY FUNCTION")][0]
    blocks, func = [], None
    for i in sheet.index[sheet.index > start]:
        nm = text.at[i, name_col]
        roman = re.match(r"^([IVX]+)\.", nm.translate(str.maketrans("ІХ", "IX")))
        if roman and nm == nm.upper() and roman.group(1) in ROMAN:
            func = ROMAN[roman.group(1)]
            blocks.append({"level": "function", "func": func, "total": None, "lines": []})
        elif nm and nm == nm.upper() and re.match(r"^[A-E]\.?\s+[A-Z]", nm):
            func = next(f for pattern, f in FUNC2_BY_HEADER if re.search(pattern, nm))
            blocks.append({"level": "sub", "func": func, "total": None, "lines": []})
        elif nm.startswith("Total expen") and blocks:
            blocks[-1]["total"] = amount[i]
        if not (para[i] and leaf[i]):
            continue
        if econ[i] not in econ2_of or amount[i] == 0:
            continue
        if YEAR == 2021 and abs(amount[i]) < 0.05:
            continue
        blocks[-1]["lines"].append((func, econ[i], round(amount[i] * 1000, 3)))
    su_lines = []
    for k, b in enumerate(blocks):
        if b["level"] == "function":
            subs = [s for s in blocks[k + 1:k + 8] if s["level"] == "sub"]
            subs = subs[:next((n for n, s in enumerate(blocks[k + 1:]) if s["level"] == "function"), len(subs))]
            if subs and b["total"] is not None and abs(b["total"] - sum(s["total"] or 0 for s in subs)) < 0.5:
                continue
        su_lines.extend(b["lines"])
    su = pd.DataFrame(su_lines, columns=["func", "econ_code", "adjusted"])
    su["executed"] = su["adjusted"]
    print(f"special units: {len(su)} lines; total {su['executed'].sum() / 1e6:,.1f} million BGN")

    # 5. Label
    clean["year"] = str(YEAR)
    for c, table in (("admin1", admin1_of), ("admin2", admin2_of), ("admin3", admin3_of)):
        clean[c] = clean["unit"].map(table).fillna("#N/A")
    for c, table in (("func1", func1_of), ("func2", func2_of), ("func3", func3_of)):
        clean[c] = clean["activity"].map(table).fillna("#N/A")
    if YEAR == 2022:
        clean.loc[clean["activity"] == "143", "func3"] = "143 n/a"
    clean["econ1"] = clean["econ_code"].map(econ1_of).fillna("#N/A")
    clean["econ2"] = clean["econ_code"].map(econ2_of).fillna("#N/A")
    code1 = np.where(clean["ibsf"] + clean["act"] == 0, 0, np.maximum(clean["ibsf"], clean["act"]))
    if YEAR == 2020:
        code1 = np.where(clean["ibsf"] == 3, 10, code1)
    clean["fin_source1"] = pd.Series(code1, index=clean.index).map(fin1_of).fillna("#N/A")
    first_digit = clean["fin_source1"].str[0].map(lambda c: fin2_of.get(int(c)) if c.isdigit() else None)
    clean["fin_source2"] = pd.Series(np.where(clean["ibsf"] == 9, clean["op"].map(fin2_of), first_digit), index=clean.index).fillna("#N/A")
    clean["exp_type"] = clean["econ2"].map(exptype_of).fillna("#N/A").replace("", "0")
    clean["transfer"] = "Excluding transfers"

    su = su.assign(year=str(YEAR), admin1="1 Central" if YEAR == 2020 else "3 Other", admin2="99 Special spending units (defense-related)",
                   admin3="9999 Special Spending Units", fin_source1="0 State budget",
                   fin_source2="00 State or municipal budget", transfer="Excluding transfers",
                   exp_type=su["econ_code"].map(exptype_of_code))
    su["func1"], su["func2"], su["func3"] = su["func"].map(func1_of), su["func"].map(func2_of), su["func"].map(func3_of)
    if YEAR == 2020:
        su["econ1"], su["econ2"] = su["econ_code"].map(exptype_of_code), su["econ_code"].map(econ1_of)
    elif YEAR == 2021:
        su["econ1"], su["econ2"] = su["econ_code"].map(econ1_of), ""
    else:
        su["econ1"], su["econ2"] = su["econ_code"].map(econ1_of), su["econ_code"].map(econ2_of)

    # 6. Assemble and write the year
    parts = [su[EXP_COLS], clean[EXP_COLS]] if YEAR == 2020 else [clean[EXP_COLS], su[EXP_COLS]]
    expenditures[YEAR] = pd.concat(parts, ignore_index=True)
    year_path = OUT_DIR / f"BGR_expenditures_{YEAR}.csv"
    expenditures[YEAR].to_csv(year_path, index=False, float_format="%.17g", lineterminator="\n", encoding="utf-8")
    print(f"wrote {year_path}: {len(expenditures[YEAR]):,} rows; executed {expenditures[YEAR]['executed'].sum() / 1e6:,.1f} million BGN")

# COMMAND ----------

# 7. Check against the delivered data: rows matched one to one on the codes that open the labels and on the
#    amounts (the code list names codes the delivered files leave as "n/a"), and, for information, on the label text.
if REFERENCE_WORKBOOK.exists():
    from openpyxl import load_workbook
    ws = load_workbook(REFERENCE_WORKBOOK, read_only=True, data_only=True)[REFERENCE_SHEET]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    delivered = pd.DataFrame([r[:len(EXP_COLS)] for r in rows if r and str(r[0]).split(".")[0].isdigit()], columns=EXP_COLS)
    delivered["year"] = delivered["year"].astype(str).str.split(".").str[0]
    for YEAR in INPUTS:
        a, b = expenditures[YEAR].copy(), delivered[delivered["year"] == str(YEAR)].copy()
        if not len(b):
            print(f"\ndelivered data {YEAR}: not in the reference workbook")
            continue
        for d in (a, b):
            for c in ("adjusted", "executed"):
                d[c] = pd.to_numeric(d[c], errors="coerce").round(2).fillna(-1)
            for c in EXP_COLS:
                if c not in ("adjusted", "executed"):
                    d[c] = d[c].fillna("").astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
                    d[c + "_code"] = d[c].str.split(" ").str[0].str.replace(".", "", regex=False)  # "98.323" and "98323" alike
            d["_n"] = d.groupby(EXP_COLS).cumcount()
            d["_n_code"] = d.groupby([c + "_code" if c not in ("adjusted", "executed") else c for c in EXP_COLS]).cumcount()
        on_codes = [c + "_code" if c not in ("adjusted", "executed") else c for c in EXP_COLS] + ["_n_code"]
        m = a.merge(b, on=on_codes, how="outer", indicator=True)
        text = a.merge(b, on=EXP_COLS + ["_n"], how="outer", indicator=True)
        print(f"\ndelivered data {YEAR}: ours {len(a):,} rows, file {len(b):,} rows; identical rows {int((m['_merge'] == 'both').sum()):,}; "
              f"only ours {int((m['_merge'] == 'left_only').sum()):,}; only file {int((m['_merge'] == 'right_only').sum()):,}"
              f" (on the label text: identical {int((text['_merge'] == 'both').sum()):,})")
        print(f"executed, million BGN: ours {a['executed'].sum() / 1e6:,.1f}, file {b['executed'].sum() / 1e6:,.1f}")
        diff = m[m["_merge"] != "both"]
        if len(diff):
            by = diff.assign(side=diff["_merge"].map({"left_only": "ours", "right_only": "file"}), paragraph=diff["econ2_code"],
                             part=np.where(diff["admin3_code"] == "9999", "special units", "extract"))
            print("rows that differ, by paragraph (executed in million BGN):")
            print(by.groupby(["part", "paragraph", "side"]).agg(rows=("executed", "size"), executed=("executed", lambda s: round(s.sum() / 1e6, 2)))
                  .unstack("side", fill_value=0).to_string())
