# Databricks notebook source
# MAGIC %run ../utils

# COMMAND ----------

# MAGIC %md
# MAGIC # Yemen BOOST expenditures: yearly Ministry of Finance files to the BOOST microdata
# MAGIC The authorities deliver one workbook per year, `YYYY.xlsx`, with a summary sheet and twelve sector sheets. Every
# MAGIC sector sheet is a matrix of spending units (rows) by economic chapters (columns). This notebook unpivots those
# MAGIC matrices into the layout of the `BOOST` tab of the delivered workbook (`Year, Func, Admin1, Econ1, econ2, Executed`)
# MAGIC and writes one `YYYY.csv` per year to `raw_microdata_csv/Yemen`. Every year whose file sits in
# MAGIC `Data from authorities/Yemen/` is processed; a new year needs no change here unless the layout changes.
# MAGIC
# MAGIC As in the `BOOST` tab: the Electricity sector is counted under *Economic Affairs*; the *Unclassified Expenditures*
# MAGIC column (chapter 99) is not carried over; a unit appears under a sector only in the years it has spending there,
# MAGIC and then with all eighteen chapters, zeros included (the tab also lists seven units with no spending at all;
# MAGIC those zero rows are not reproduced). The last two cells check the parse against each file's own
# MAGIC summary sheet and against the `BOOST` tab of the delivered workbook (see `BOOST_tab_stacking_check.md`).

# COMMAND ----------

import pandas as pd
from openpyxl import load_workbook

COUNTRY = 'Yemen'
OUT_DIR = Path(prepare_raw_microdata_csv_dir(COUNTRY))
BASE = Path(f"{RAW_INPUT_DIR}/{COUNTRY}")
REFERENCE_SHEET = "BOOST"  # the microdata tab of the delivered workbook (input_excel_filename), for the check cell

YEARS = sorted(int(f.stem) for f in BASE.glob("2???.xlsx"))
print("years:", ", ".join(map(str, YEARS)))

# Sector sheet -> Func, labelled as in the BOOST tab. Sheet names are as Excel truncated them (31 characters);
# the Public Order sheet's name ends in a space. The Electricity sector is counted under Economic Affairs.
SECTOR_FUNC = {
    'Public Services Sector': 'General Public Services',
    'Education sector': 'Education',
    'Public Order and Public Safety ': 'Public Safety',
    'Health sector': 'Health',
    'Entertainment, Culture and Reli': 'Culture and arts',
    'Social Protection Sector': 'Social Protection',
    'Economic Affairs Sector': 'Economic Affairs',
    'Environmental Protection Sector': 'Environment',
    'Defense sector': 'Defense',
    'Population and Community Facili': 'Population and Community Development',
    'Electricity sector': 'Economic Affairs',
    'Public debt': 'Debt',
}
SUMMARY_SHEET = 'القطاعات'  # "Sectors": chapter x sector totals, used by the check cell

# Chapter code (row 8 of every sector sheet) -> (Econ1, econ2), labelled as in the BOOST tab. The units digit of
# the code is the Title (I wages, II goods and services, III transfers, IV non-financial assets, V financial
# assets and liabilities); interest (32) is the one chapter of Title II that the BOOST tab gives its own Econ1.
CHAPTERS = {
    11: ('Wages and salaries', 'Salaries, Wages and Equivalent'),
    21: ('Wages and salaries', 'Social Contributions'),
    12: ('Use of goods and services', 'Goods and services'),
    22: ('Use of goods and services', 'Maintenance'),
    32: ('Interest', 'Interest'),
    42: ('Use of goods and services', 'Depreciation of Fixed Capital'),
    52: ('Use of goods and services', 'Expenditures on Property Other Than Interest'),
    13: ('Grants and transfers', 'Financial Subsidies'),
    23: ('Grants and transfers', 'Grants'),
    33: ('Grants and transfers', 'Social Benefits'),
    43: ('Grants and transfers', 'Other Financial Transfers and Subsidies'),
    14: ('Acquisition on non financial assets', 'Acquisition of Fixed Assets'),
    24: ('Acquisition on non financial assets', 'Acquisition of Inventories'),
    34: ('Acquisition on non financial assets', 'Acquisition of Non‑productive Assets'),
    15: ('Assets and liabilities', 'Local Lending and Acquisition of Domestic Financial Assets'),
    25: ('Assets and liabilities', 'Local Lending and Acquisition of Foreing Financial Assets'),
    35: ('Assets and liabilities', 'Repayment of Domestic Loans and Redemption of Domestic Securities (Excluding Equity)'),
    45: ('Assets and liabilities', 'Repayment of Foreign Loans and Redemption of Foreign Securities (Excluding Equity)'),
}
UNCLASSIFIED = 99  # "Unclassified Expenditures" (column AB): parsed for the summary check, not written

# Sector sheet layout: the year in B3; chapter codes in row 8; spending units from row 11, Arabic name in
# column C and English name in column D, down to the row whose English name is "Total".
YEAR_CELL = (3, 2)
CODE_ROW = 8
FIRST_UNIT_ROW = 11
UNIT_AR_COL, UNIT_EN_COL = 3, 4
TOTAL_LABEL = 'Total'

OUT_COLS = ['Year', 'Func', 'Admin1', 'Econ1', 'econ2', 'Executed']

# COMMAND ----------

def read_sector(ws, year):
    """One sector sheet -> rows (unit, chapter code, amount), every chapter including 99, every unit."""
    rows = list(ws.iter_rows(values_only=True))
    sheet_year = rows[YEAR_CELL[0] - 1][YEAR_CELL[1] - 1]
    assert sheet_year == year, f"{year} '{ws.title}': the sheet says year {sheet_year!r}"
    codes = {j: int(v) for j, v in enumerate(rows[CODE_ROW - 1]) if isinstance(v, (int, float))}
    assert set(codes.values()) == set(CHAPTERS) | {UNCLASSIFIED}, f"{year} '{ws.title}': chapter codes {sorted(codes.values())}"
    recs = []
    for i, r in enumerate(rows[FIRST_UNIT_ROW - 1:], start=FIRST_UNIT_ROW):
        unit = r[UNIT_EN_COL - 1]
        if unit == TOTAL_LABEL:
            break
        assert unit is not None, f"{year} '{ws.title}' row {i}: no unit name"
        for j, code in codes.items():
            assert r[j] is None or isinstance(r[j], (int, float)), f"{year} '{ws.title}' row {i} col {j + 1}: {r[j]!r}"
            recs.append((ws.title, unit, code, r[j]))
    else:
        raise AssertionError(f"{year} '{ws.title}': no '{TOTAL_LABEL}' row")
    return recs

parsed, written = {}, {}
for YEAR in YEARS:
    wb = load_workbook(BASE / f"{YEAR}.xlsx", read_only=True, data_only=True)
    assert set(SECTOR_FUNC) <= set(wb.sheetnames), f"{YEAR}: sector sheets missing {set(SECTOR_FUNC) - set(wb.sheetnames)}"
    long = pd.DataFrame([rec for name in SECTOR_FUNC for rec in read_sector(wb[name], YEAR)],
                        columns=['sector', 'Admin1', 'code', 'Executed'])
    wb.close()
    parsed[YEAR] = long
    blanks = int(long['Executed'].isna().sum())

    # The BOOST tab layout: one row per (Func, unit, chapter); the Electricity sector adds to the Economic Affairs
    # rows of the same unit; a unit stays under a Func only if it has some spending there in the year.
    long = long[long['code'] != UNCLASSIFIED]
    df = (long.assign(Func=long['sector'].map(SECTOR_FUNC))
              .groupby(['Func', 'Admin1', 'code'], sort=False)['Executed'].sum(min_count=1)
              .reset_index())
    has_spending = df.groupby(['Func', 'Admin1'])['Executed'].transform(lambda s: s.fillna(0).ne(0).any())
    df = df[has_spending].reset_index(drop=True)
    df['Year'] = YEAR
    df['Econ1'] = df['code'].map(lambda c: CHAPTERS[c][0])
    df['econ2'] = df['code'].map(lambda c: CHAPTERS[c][1])
    written[YEAR] = df[OUT_COLS]

    year_path = OUT_DIR / f"{YEAR}.csv"
    written[YEAR].to_csv(year_path, index=False, float_format="%.17g", lineterminator="\n", encoding="utf-8")
    print(f"{YEAR}: {len(long) // len(CHAPTERS):,} unit x sector rows read ({blanks} blank cells); wrote {year_path}: "
          f"{len(written[YEAR]):,} rows, {written[YEAR][['Func', 'Admin1']].drop_duplicates().shape[0]} unit x Func pairs; "
          f"executed {written[YEAR]['Executed'].sum():,.2f} (unclassified, not written: {parsed[YEAR].loc[parsed[YEAR]['code'] == UNCLASSIFIED, 'Executed'].sum():,.2f})")

