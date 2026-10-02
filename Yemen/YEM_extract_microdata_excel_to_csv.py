# Databricks notebook source
# MAGIC %run ../utils

# COMMAND ----------

# Yemen BOOST microdata extraction (test input for the formula -> Python mapping).
#
# Dumps the workbook's `BOOST` tab to CSV with no transformation. The tab is the microdata that the
# `Executed` sheet aggregates via SUMIFS over its named column ranges (Year, Func, Admin1, Econ1, econ2,
# Executed), so a pipeline run on this extract must reproduce every Executed cell exactly; the one in
# verification.md does (75 / 75 cells, 2023-2025). The production input is the rebuild of the yearly
# sector files (YEM_extract_raw_microdata_excel_to_csv.py, same six columns), because the tab's 2025
# block is mis-stacked (BOOST_tab_stacking_check.md).
#
# One file per year, YYYY.csv, in the same layout as the rebuild, so that YEM_transform_load_dlt.py runs
# on this folder unchanged: point its COUNTRY_MICRODATA_DIR at microdata_csv/Yemen instead of
# raw_microdata_csv/Yemen.
#
# Minimal hygiene only: the six named columns, blank rows dropped (the tab's used range runs 63,000
# rows past the data and holds one blank row inside it), Year as an integer. Labels are kept verbatim,
# spaces and all, and the tab's six duplicated zero rows of 2025 are kept: the SUMIFS see them too.

import pandas as pd
from openpyxl import load_workbook

COUNTRY = 'Yemen'
RAW_SHEET = 'BOOST'
EXECUTED_SHEET = 'Executed'
COLUMNS = ['Year', 'Func', 'Admin1', 'Econ1', 'econ2', 'Executed']

microdata_csv_dir = prepare_microdata_csv_dir(COUNTRY)
filename = input_excel_filename(COUNTRY)

df = pd.read_excel(filename, sheet_name=RAW_SHEET, header=0, usecols=COLUMNS)
df = df.dropna(how='all')
df['Year'] = df['Year'].astype(int)

for year, part in df.groupby('Year', sort=True):
    csv_file_path = f'{microdata_csv_dir}/{year}.csv'
    part.to_csv(csv_file_path, index=False, float_format='%.17g', lineterminator='\n', encoding='utf-8')
    print(f'wrote {len(part):,} rows -> {csv_file_path}')

# COMMAND ----------

# Check: the extract against the Executed sheet's own total, EXP_ECON_TOT_EXP_EXE (row 2):
# SUM(SUMIFS(Executed, Year, <year>, Econ1, "<>assets*")), read as the cached cell value.
ws = load_workbook(filename, read_only=True, data_only=True)[EXECUTED_SHEET]
rows = ws.iter_rows(min_row=1, max_row=2, values_only=True)
years, total = next(rows), next(rows)
assert total[0] == 'EXP_ECON_TOT_EXP_EXE', total[0]
cached = {int(y): v for y, v in zip(years[2:], total[2:]) if isinstance(y, (int, float)) and isinstance(v, (int, float))}
ours = df[~df['Econ1'].str.lower().str.startswith('assets', na=False)].groupby('Year')['Executed'].sum()
for year, value in cached.items():
    diff = ours.get(year, 0.0) - value
    print(f'{year}: extract {ours.get(year, 0.0):,.2f}  Executed!row 2 {value:,.2f}  diff {diff:,.2f}  {"ok" if abs(diff) < 0.5 else "MISMATCH"}')
